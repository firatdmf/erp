"""Excel (.xlsx) export of an invoice document (accounting/invoice_doc.py) —
a bordered, document-style sheet that mirrors the printed invoice: issuer,
customer or supplier, lines and total.
"""
from io import BytesIO

from django.http import HttpResponse

from erp.xlsx_utils import (
    cell, merge, merge_border, section, kv_full, kv_pair,
    GRID, RULE, FILL_HEAD, RIGHT, LEFT, TOP, TEXT,
    F_TITLE, F_DOCNO, F_HEAD, F_VAL, F_VALB, F_TOTAL,
)


def _dec(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _date(v):
    try:
        return v.strftime("%d %b %Y") if v else "—"
    except Exception:
        return "—"


def build_invoice_workbook(doc):
    """`doc` is an invoice_doc.InvoiceDoc — the same object the printed
    page renders, so the two can never state different figures."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    ccode = doc.currency_code
    money = (f'#,##0.00" {ccode}"' if ccode else "#,##0.00")
    iss, party = doc.issuer, doc.party
    lines = doc.lines

    icols = [("Description", 40, "desc"), ("SKU", 18, "sku"),
             ("Qty", 12, "qty"), ("Unit", 7, "unit"),
             ("Unit Price", 14, "price"), ("Line Total", 16, "total")]
    NCOLS = len(icols)
    mid = (NCOLS + 1) // 2

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.sheet_view.showGridLines = False
    for i, (h, w, k) in enumerate(icols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    title = "PURCHASE INVOICE" if doc.kind == "purchase" else "INVOICE"

    # ── Header ──
    r = 1
    cell(ws, r, 1, (iss.name or "").upper(), font=F_TITLE); merge(ws, r, 1, mid)
    cell(ws, r, mid + 1, title, font=F_DOCNO, align=RIGHT); merge(ws, r, mid + 1, NCOLS)
    r += 1
    cell(ws, r, mid + 1, f"No: {doc.number or '—'}", font=F_VALB, align=RIGHT); merge(ws, r, mid + 1, NCOLS)
    for c in range(1, NCOLS + 1):
        ws.cell(r, c).border = RULE
    r += 2

    # ── Invoice details ──
    r = section(ws, r, "INVOICE DETAILS", NCOLS)
    r = kv_pair(ws, r, "Invoice No", doc.number or "—", "Date", _date(doc.date), NCOLS)
    r = kv_pair(ws, r, "Currency", ccode or "—", "Lines", str(len(lines)), NCOLS)
    r += 1

    def party_block(r, heading, p):
        r = section(ws, r, heading, NCOLS)
        r = kv_full(ws, r, "Name", p.name or "—", NCOLS, bold_value=True)
        if p.address:
            r = kv_full(ws, r, "Address", p.address, NCOLS)
        loc = ", ".join([x for x in [p.city, p.country] if x])
        if loc:
            r = kv_full(ws, r, "City / Country", loc, NCOLS)
        if p.phone:
            r = kv_full(ws, r, "Phone", p.phone, NCOLS)
        if p.fax:
            r = kv_full(ws, r, "Fax", p.fax, NCOLS)
        if p.email:
            r = kv_full(ws, r, "Email", p.email, NCOLS)
        if p.tax_office or p.tax_number:
            r = kv_full(ws, r, "Tax Office / No",
                        f"{p.tax_office}  ·  {p.tax_number}".strip(" ·"), NCOLS)
        return r + 1

    r = party_block(r, "ISSUER", iss)
    r = party_block(r, "SUPPLIER" if doc.kind == "purchase" else "BILL TO", party)

    # ── Items ──
    r = section(ws, r, f"ITEMS ({len(lines)})", NCOLS)
    for i, (h, w, k) in enumerate(icols, 1):
        cell(ws, r, i, h, font=F_HEAD, fill=FILL_HEAD, border=GRID,
             align=(LEFT if k in ("desc", "sku", "unit") else RIGHT))
    r += 1

    for ln in lines:
        vals = {
            "desc": (ln.description or "—", None, TOP),
            "sku": (ln.sku or "—", TEXT, LEFT),
            "qty": (_dec(ln.quantity), "#,##0.00", RIGHT),
            "unit": (ln.unit or "", TEXT, LEFT),
            "price": (_dec(ln.unit_price), money, RIGHT),
            "total": (_dec(ln.total), money, RIGHT),
        }
        for i, (h, w, k) in enumerate(icols, 1):
            v, fmt, al = vals[k]
            cell(ws, r, i, v, font=F_VAL, border=GRID, align=al, fmt=fmt)
        r += 1

    # ── Total ──
    r += 1
    lc1, lc2, vc = NCOLS - 2, NCOLS - 1, NCOLS
    cell(ws, r, lc1, "TOTAL", font=F_TOTAL, border=GRID, align=RIGHT)
    merge(ws, r, lc1, lc2)
    merge_border(ws, r, lc1, lc2, GRID)
    cell(ws, r, vc, _dec(doc.total), font=F_TOTAL, border=GRID, align=RIGHT, fmt=money)
    r += 1

    # ── Notes ──
    if doc.notes:
        r += 1
        r = section(ws, r, "NOTES", NCOLS)
        cell(ws, r, 1, doc.notes, font=F_VAL, border=GRID, align=TOP)
        merge(ws, r, 1, NCOLS)
        merge_border(ws, r, 1, NCOLS, GRID)
        ws.row_dimensions[r].height = 46

    return wb


def workbook_response(doc):
    """The document as an .xlsx download."""
    wb = build_invoice_workbook(doc)
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{doc.filename}.xlsx"'
    return resp
