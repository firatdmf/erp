"""The "Created with Nejum" credit on the documents a customer receives.

On by default for the house, off with settings.NEJUM_CREDIT — and when it
is off, every document drops it. Invoices never carry it.
"""
import io
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse

from accounting.models import Book, CurrencyCategory
from accounting.models_accounts import CurrentAccount
from marketing.models import Product
from operating.models import (Order, OrderItem, OrderStockReservation, Pack, Warehouse,
                              WarehouseProduct, WarehouseProductItem)

LINE = "Created with Nejum"


class NejumCreditBase(TestCase):
    @patch("marketing.utils.bunny_storage.upload_to_bunny")
    def setUp(self, mock_upload):
        mock_upload.return_value = "https://mock-cdn.net/qr.png"
        self.book = Book.objects.create(name="Laleli Fabric")
        usd = CurrencyCategory.objects.create(code="USD", name="US Dollar", symbol="$")
        account = CurrentAccount.objects.create(
            book=self.book, code="C-601", name="Oleg", type="customer",
            default_currency=usd)
        self.order = Order.objects.create(
            order_number="DK0000601", current_account=account,
            notify_customer=True)
        product = Product.objects.create(title="Velvet Moss", sku="KZL000601", price=10)
        item = OrderItem.objects.create(order=self.order, product=product,
                                        quantity=Decimal("40.00"), price=Decimal("3.25"))
        wh = Warehouse.objects.create(name="WH", accounting_book=self.book)
        wp = WarehouseProduct.objects.create(warehouse=wh, name="Velvet Moss",
                                             sku="KZL000601", quantity=50)
        roll = WarehouseProductItem.objects.create(
            product=wp, quantity=Decimal("50"), quantity_remaining=Decimal("50"),
            barcode="LZK0300003477")
        pack = Pack.objects.create(order=self.order, pack_number=1)
        OrderStockReservation.objects.create(
            order=self.order, order_item=item, stock_item=roll,
            warehouse_product=wp, quantity=Decimal("40.00"), pack=pack)

        user = User.objects.create_superuser("boss_credit", "b@t.com", "pw")
        self.client.force_login(user)

    def pdf_strings(self, url):
        """Every string drawn straight onto the PDF canvas (the footer
        hook draws there; the table content goes through Paragraphs)."""
        from reportlab.pdfgen.canvas import Canvas
        drawn = []
        real = Canvas.drawString

        def spy(canvas, x, y, text, *a, **kw):
            drawn.append(text)
            return real(canvas, x, y, text, *a, **kw)

        with patch.object(Canvas, "drawString", spy):
            resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        return drawn

    def xlsx_values(self, url):
        """Every cell value, plus the printed page footer."""
        import openpyxl
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        return self.sheet_values(ws)

    @staticmethod
    def sheet_values(ws):
        values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        return values + [ws.oddFooter.center.text]

    def sent_email(self):
        from operating.order_notifications import send_order_event_email
        with patch("operating.order_notifications._resolve_customer_email",
                   return_value="buyer@example.com"), \
             patch("operating.order_notifications._send_via_gmail_oauth",
                   return_value=True) as send:
            self.assertTrue(send_order_event_email(self.order, "shipped", attach_pdf=False))
        args, kwargs = send.call_args
        return args[2], kwargs["text_body"]


class TheCreditIsOn(NejumCreditBase):

    def test_the_printed_order(self):
        """On the footer line, beside who sent it and when."""
        import re
        resp = self.client.get(reverse("operating:order_print", args=[self.order.pk]))
        self.assertContains(resp, LINE)
        self.assertContains(resp, 'href="https://nejum.com"')
        foot = re.search(r'<table class="foot".*?</table>', resp.content.decode(), re.S)
        self.assertIsNotNone(foot)
        self.assertIn(LINE, foot.group(0))

    def test_the_packing_list_pdf(self):
        self.assertIn(LINE, self.pdf_strings(
            reverse("operating:order_packing_list_pdf", args=[self.order.pk])))

    def test_the_packing_list_excel(self):
        self.assertIn(LINE, self.xlsx_values(
            reverse("operating:export_packing_list_excel", args=[self.order.pk])))

    def test_the_order_excel(self):
        self.assertIn(LINE, self.xlsx_values(
            reverse("operating:order_excel", args=[self.order.pk])))

    def test_the_combined_order_excel(self):
        from operating.order_excel import build_combined_workbook
        ws = build_combined_workbook([self.order]).active
        self.assertIn(LINE, self.sheet_values(ws))

    def test_excel_carries_it_in_the_footer_not_a_cell(self):
        """People add rows under the totals; ours must not be one."""
        import openpyxl
        resp = self.client.get(reverse("operating:order_excel", args=[self.order.pk]))
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        cells = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertNotIn(LINE, cells)
        self.assertEqual(ws.oddFooter.center.text, LINE)

    def test_the_order_email(self):
        html, text = self.sent_email()
        self.assertIn(LINE, html)
        self.assertIn(LINE, text)

    def test_the_pdf_the_email_attaches(self):
        """It has a footer line of its own, so the credit joins it."""
        from reportlab.platypus import Table
        from operating.order_notifications import _render_order_pdf
        tables = []
        real = Table.__init__

        def spy(table, data, *a, **kw):
            tables.append(data)
            return real(table, data, *a, **kw)

        with patch.object(Table, "__init__", spy):
            name, data = _render_order_pdf(self.order)
        self.assertTrue(data.startswith(b"%PDF"))
        texts = [cell.getPlainText() for rows in tables for row in rows for cell in row
                 if hasattr(cell, "getPlainText")]
        self.assertIn(LINE, texts)

    def test_the_invoice_never_carries_it(self):
        resp = self.client.get(reverse("operating:order_invoice", args=[self.order.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, "Nejum ·")
        self.assertNotIn(LINE, self.xlsx_values(
            reverse("operating:order_invoice_excel", args=[self.order.pk])))

    def test_it_follows_the_language(self):
        from django.utils import translation
        from erp.nejum_credit import nejum_credit
        with translation.override("tr"):
            self.assertEqual(nejum_credit(), "Nejum ile düzenlenmiştir")


@override_settings(NEJUM_CREDIT=False)
class TheCreditIsOff(NejumCreditBase):

    def test_no_document_carries_it(self):
        resp = self.client.get(reverse("operating:order_print", args=[self.order.pk]))
        self.assertNotContains(resp, LINE)
        self.assertNotIn(LINE, self.pdf_strings(
            reverse("operating:order_packing_list_pdf", args=[self.order.pk])))
        self.assertNotIn(LINE, self.xlsx_values(
            reverse("operating:export_packing_list_excel", args=[self.order.pk])))
        self.assertNotIn(LINE, self.xlsx_values(
            reverse("operating:order_excel", args=[self.order.pk])))
        html, text = self.sent_email()
        self.assertNotIn(LINE, html)
        self.assertNotIn(LINE, text)
