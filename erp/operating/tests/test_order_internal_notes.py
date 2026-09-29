"""An order carries two notes: the one printed for the customer, and one
for the team that never leaves the order page.

The printed notes go out on the order PDF, the emailed PDF, the Excel and
the sales invoice, so anything said about cost, margin or a stock
correction has nowhere to live but internal_notes.

Run with:
    python manage.py test operating.tests.test_order_internal_notes
"""
import io
import json
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from accounting.invoice_doc import build_order_doc
from accounting.models import Book, CurrencyCategory
from crm.models import Contact
from marketing.models import Product
from operating.models import Order, OrderChange

PRINTED = "Delivery in 3 weeks."
INTERNAL = "Everything at cost: $9,973.76"


class InternalNotesTest(TestCase):
    @patch("marketing.utils.bunny_storage.upload_to_bunny")
    def setUp(self, mock_upload):
        mock_upload.return_value = "https://mock-cdn.net/qr.png"
        CurrencyCategory.objects.create(code="USD", name="US Dollar", symbol="$")
        self.book = Book.objects.create(name="Ergene Fabric")
        self.customer = Contact.objects.create(name="Hanefi")
        Product.objects.create(title="Krep", sku="KRP", featured=False)
        user = get_user_model().objects.create_superuser("seller_notes", "s@t.com", "pw")
        user.member.books.add(self.book)
        user.member.default_book = self.book
        user.member.save()
        self.client.force_login(user)

    def _post(self, url=None, item_id=None, **fields):
        line = {
            "item_no": 1, "product": {"sku": "KRP", "variant": False},
            "description": "", "quantity": 100, "outsourced": 0,
            "price": 2, "is_custom_curtain": False, "rolls": [],
        }
        if item_id:
            line["item_id"] = item_id
        data = {
            "customer_type": "contact", "customer_pk": self.customer.pk,
            "book": self.book.pk,
            "product_json_input": json.dumps([line]),
        }
        data.update(fields)
        return self.client.post(url or reverse("operating:create_order"), data)

    def _order(self):
        self._post(notes=PRINTED, internal_notes=INTERNAL)
        return Order.objects.get()

    def test_the_form_saves_both(self):
        order = self._order()
        self.assertEqual(order.notes, PRINTED)
        self.assertEqual(order.internal_notes, INTERNAL)
        self.assertEqual(order.original_snapshot["internal_notes"], INTERNAL)

    def test_the_edit_form_reopens_with_both(self):
        """The edit view saves whatever the boxes hold, so a box that
        opened empty wiped the notes on every edit."""
        order = self._order()
        html = self.client.get(reverse("operating:edit_order", args=[order.pk])).content.decode()
        self.assertIn(f">{PRINTED}</textarea>", html)
        self.assertIn(f">{INTERNAL}</textarea>", html)

    def test_editing_keeps_what_the_boxes_hold(self):
        order = self._order()
        self._post(url=reverse("operating:edit_order", args=[order.pk]),
                   item_id=order.items.get().pk, notes=PRINTED, internal_notes="Rechecked")
        order.refresh_from_db()
        self.assertEqual(order.notes, PRINTED)
        self.assertEqual(order.internal_notes, "Rechecked")

    def test_the_order_page_edits_them_on_their_own(self):
        order = self._order()
        page = self.client.get(reverse("operating:order_detail", args=[order.pk]))
        self.assertContains(page, "od-internal-notes-edit")
        self.assertContains(page, INTERNAL)

        r = self.client.post(reverse("operating:order_detail", args=[order.pk]),
                             {"action": "update_internal_notes", "internal_notes": "Call first"})
        self.assertEqual(r.json(), {"ok": True})
        order.refresh_from_db()
        self.assertEqual(order.internal_notes, "Call first")
        self.assertEqual(order.notes, PRINTED)
        self.assertTrue(OrderChange.objects.filter(
            order=order, field="internal_notes", new_value="Call first").exists())

    def test_nothing_the_customer_gets_carries_them(self):
        order = self._order()
        printout = self.client.get(reverse("operating:order_print", args=[order.pk]), {"html": "1"})
        self.assertContains(printout, PRINTED)
        self.assertNotContains(printout, INTERNAL)

        self.assertEqual(build_order_doc(order).notes, PRINTED)

        from openpyxl import load_workbook
        sheet = self.client.get(reverse("operating:order_excel", args=[order.pk]))
        cells = " ".join(str(c.value) for ws in load_workbook(io.BytesIO(sheet.content))
                         for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertIn(PRINTED, cells)
        self.assertNotIn(INTERNAL, cells)
