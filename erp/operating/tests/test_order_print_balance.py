"""The customer's balance under an order sheet's total.

    Previous balance    $100.00
    This order          $390.00
    Balance due         $490.00

"Previous" leaves out the order's own sale rows, so the sheet reads the
same before the order ships (no sale yet) and after (sale posted). Always
printed, "In credit" when we owe the customer, nothing at zero.

Run:
    python manage.py test operating.tests.test_order_print_balance
"""
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.test import TestCase
from django.urls import reverse

from accounting.models import Book, CurrencyCategory
from accounting.models_accounts import CurrentAccount, CurrentAccountMovement
from crm.models import Contact
from marketing.models import Product, ProductVariant

from operating.models import Order, OrderItem

User = get_user_model()


class OrderPrintBalance(TestCase):

    @patch("marketing.utils.bunny_storage.upload_to_bunny")
    def setUp(self, mock_upload):
        mock_upload.return_value = "https://mock-cdn.net/qr.png"
        self.usd = CurrencyCategory.objects.create(code="USD", name="US Dollar", symbol="$")
        self.laleli = Book.objects.create(name="Laleli Fabric")
        self.ergene = Book.objects.create(name="Ergene Fabric")
        self.oleg = Contact.objects.create(name="OLEG MOTUZENKO")
        self.crepe = Product.objects.create(title="Crepe", sku="KZL000315", price=10)
        self.black = ProductVariant.objects.create(product=self.crepe, variant_sku="KZL000315.BLACK")
        # 156.00 × 2.50 = 390.00
        self.order = self._order(self.laleli, self.oleg, "DK-284", Decimal("156.00"), Decimal("2.50"))
        self.account = self.order.current_account
        self.boss = User.objects.create_superuser("boss", "b@t.com", "pw")
        self.client.force_login(self.boss)

    def _order(self, book, contact, number, qty, price):
        account = None
        if contact:
            account, _ = CurrentAccount.objects.get_or_create(
                book=book, contact=contact,
                defaults=dict(code=f"C-{book.pk}-{contact.pk}", name=contact.name,
                              type="customer", default_currency=self.usd))
        order = Order.objects.create(order_number=number, current_account=account, contact=contact)
        OrderItem.objects.create(order=order, product=self.crepe, product_variant=self.black,
                                 quantity=qty, price=price)
        return order

    def _move(self, amount, movement_type="opening", order=None, account=None):
        account = account or self.account
        source = {}
        if order is not None:
            source = dict(source_type=ContentType.objects.get_for_model(Order), source_id=order.pk)
        CurrentAccountMovement.objects.create(
            current_account=account, book=account.book, date=date(2026, 10, 1),
            amount=Decimal(amount), currency=self.usd, movement_type=movement_type, **source)
        account.recompute_balance(save=True)

    def _sheet(self, order=None):
        order = order or self.order
        return self.client.get(reverse("operating:order_print", args=[order.pk]), {"html": "1"})

    def _only_balance(self, resp):
        balances = resp.context["order_balances"]
        self.assertEqual(len(balances), 1)
        return balances[0]

    def test_an_open_order_adds_to_what_was_owed(self):
        self._move("100.00")
        resp = self._sheet()
        bal = self._only_balance(resp)
        self.assertEqual((bal["previous"], bal["this_orders"], bal["after"]),
                         (Decimal("100.00"), Decimal("390.00"), Decimal("490.00")))
        self.assertContains(resp, "Previous balance")
        self.assertContains(resp, "Balance due")
        self.assertContains(resp, "$490.00")

    def test_a_shipped_order_is_not_counted_twice(self):
        """Once it ships its sale is on the account; "previous" must not
        carry it as well, or the sheet would bill it twice."""
        self._move("100.00")
        self._move("390.00", "order_sale", order=self.order)
        bal = self._only_balance(self._sheet())
        self.assertEqual((bal["previous"], bal["after"]), (Decimal("100.00"), Decimal("490.00")))

    def test_a_payment_already_made_is_counted(self):
        self._move("100.00")
        self._move("-200.00", "collection", order=self.order)
        bal = self._only_balance(self._sheet())
        self.assertEqual((bal["previous"], bal["after"]), (Decimal("-100.00"), Decimal("290.00")))

    def test_when_we_owe_them_it_says_in_credit(self):
        self._move("-500.00", "collection")
        resp = self._sheet()
        bal = self._only_balance(resp)
        self.assertEqual(bal["after"], Decimal("-110.00"))
        self.assertContains(resp, "In credit")
        self.assertContains(resp, "$110.00")
        self.assertContains(resp, "$500.00 in credit")

    def test_nothing_prints_when_it_comes_to_zero(self):
        self._move("-390.00", "collection")
        resp = self._sheet()
        self.assertEqual(resp.context["order_balances"], [])
        self.assertNotContains(resp, "Previous balance")

    def test_an_order_with_no_customer_prints_no_balance(self):
        """A walk-in sale posts to the shared retail account — the shop's
        balance, not the customer's."""
        walk_in = self._order(self.laleli, None, "R-1", Decimal("1"), Decimal("5"))
        walk_in.current_account = self.account
        walk_in.save()
        self._move("100.00")
        self.assertEqual(self._sheet(walk_in).context["order_balances"], [])

    def test_a_sheet_across_books_shows_each_account(self):
        self._move("100.00")
        other = self._order(self.ergene, self.oleg, "DK-291", Decimal("19.50"), Decimal("4.50"))
        self._move("50.00", account=other.current_account)
        resp = self.client.get(reverse("operating:order_print_combined"),
                               {"ids": f"{self.order.pk},{other.pk}", "html": "1"})
        after = {b["account"].book.name: b["after"] for b in resp.context["order_balances"]}
        self.assertEqual(after, {"Laleli Fabric": Decimal("490.00"),
                                 "Ergene Fabric": Decimal("137.75")})
        self.assertContains(resp, "These orders")

    def test_the_excel_carries_the_same_figures(self):
        from openpyxl import load_workbook
        self._move("100.00")
        resp = self.client.get(reverse("operating:order_excel", args=[self.order.pk]))
        ws = load_workbook(BytesIO(resp.content)).active
        values = [[c for c in row if c is not None] for row in ws.iter_rows(values_only=True)]
        self.assertIn(["Previous balance", 100.0], values)
        self.assertIn(["This order", 390.0], values)
        self.assertIn(["Balance due", 490.0], values)

    def test_the_combined_excel_carries_each_account(self):
        from openpyxl import load_workbook
        self._move("-500.00", "collection")
        other = self._order(self.ergene, self.oleg, "DK-291", Decimal("19.50"), Decimal("4.50"))
        resp = self.client.get(reverse("operating:order_excel_combined"),
                               {"ids": f"{self.order.pk},{other.pk}"})
        ws = load_workbook(BytesIO(resp.content)).active
        values = [[c for c in row if c is not None] for row in ws.iter_rows(values_only=True)]
        self.assertIn(["Previous balance (in credit)", 500.0], values)
        self.assertIn(["In credit", 110.0], values)
        self.assertIn(["Balance due", 87.75], values)
