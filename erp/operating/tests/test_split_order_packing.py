"""An order split across books is packed as the one shipment it is.

Two orders in the books, one set of sacks: every half's rolls go into the
same packages, numbered as one run, and either half opens the same packing
list. A sack may hold goods off both books' shelves.

Run with:
    python manage.py test operating.tests.test_split_order_packing
"""
import uuid
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from accounting.models import Book, CurrencyCategory
from accounting.models_accounts import CurrentAccount
from authentication.models import Permission
from crm.models import Contact
from marketing.models import Product, ProductCategory
from operating.models import (Order, OrderItem, OrderStockReservation, Pack, Warehouse,
                              WarehouseProduct, WarehouseProductItem)

CDN = "https://cdn.test/qr.png"


@patch("marketing.utils.bunny_storage.upload_to_bunny", return_value=CDN)
class SplitOrderPacking(TestCase):
    @patch("marketing.utils.bunny_storage.upload_to_bunny", return_value=CDN)
    def setUp(self, _upload):
        usd = CurrencyCategory.objects.create(code="USD", name="US Dollar", symbol="$")
        customer = Contact.objects.create(name="Anna Lugansk")
        group = uuid.uuid4()
        self.laleli = Book.objects.create(name="Laleli Fabric")
        self.ergene = Book.objects.create(name="Ergene Fabric")
        self.laleli_half, self.laleli_roll = self._half(
            self.laleli, "ORD-2026-000023", customer, group, usd, "N1408T.G23", "LAL1")
        self.ergene_half, self.ergene_roll = self._half(
            self.ergene, "ORD-2026-000024", customer, group, usd, "N1477.G06", "ERG1")
        user = User.objects.create_user("firat", password="pw")
        user.member.books.add(self.laleli, self.ergene)
        user.member.permissions.add(Permission.objects.get_or_create(name="print_labels")[0])
        self.client.force_login(user)

    def _half(self, book, number, customer, group, currency, sku, barcode):
        account = CurrentAccount.objects.create(
            book=book, code=f"C-{book.pk}", name=f"Anna ({book.name})", type="customer",
            contact=customer, default_currency=currency)
        order = Order.objects.create(order_number=number, current_account=account,
                                     contact=customer, currency=currency, split_group=group)
        product = Product.objects.create(
            title=sku, sku=sku, price=1,
            category=ProductCategory.objects.get_or_create(name="fabric")[0])
        item = OrderItem.objects.create(order=order, product=product,
                                        quantity=Decimal("30"), price=Decimal("2"))
        wp = WarehouseProduct.objects.create(
            warehouse=Warehouse.objects.create(name=book.name, accounting_book=book),
            name=sku, sku=sku, quantity=Decimal("30"))
        roll = WarehouseProductItem.objects.create(
            product=wp, quantity=Decimal("30"), quantity_remaining=Decimal("30"), barcode=barcode)
        reservation = OrderStockReservation.objects.create(
            order=order, order_item=item, stock_item=roll, warehouse_product=wp,
            quantity=Decimal("30"))
        return order, reservation

    def _crud(self, order, **data):
        return self.client.post(reverse("operating:order_packing_list", args=[order.pk]), data)

    def test_either_half_opens_one_screen_with_both_halves_rolls(self, _upload):
        page = self.client.get(reverse("operating:order_pack_scan", args=[self.ergene_half.pk]))
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context["number"], "ORD-2026-000023 & ORD-2026-000024")
        self.assertEqual({r.pk for r in page.context["reservations"]},
                         {self.laleli_roll.pk, self.ergene_roll.pk})
        # The first package is made once, on the first half, for both.
        self.assertEqual([(p.order_id, p.pack_number) for p in page.context["packs"]],
                         [(self.laleli_half.pk, 1)])
        self.client.get(reverse("operating:order_pack_scan", args=[self.laleli_half.pk]))
        self.assertEqual(Pack.objects.filter(order__split_group=self.laleli_half.split_group).count(), 1)

    def test_one_sack_takes_both_halves_rolls(self, _upload):
        sack = Pack.objects.create(order=self.laleli_half, pack_number=1)
        # Dragged on the Ergene half's screen, and scanned on the Laleli
        # half's: neither cares which half the roll or the sack is on.
        r = self.client.post(reverse("operating:order_pack_reserve_assign_pack",
                                     args=[self.ergene_half.pk]),
                             {"reservation_id": self.laleli_roll.pk, "pack_id": sack.pk})
        self.assertTrue(r.json()["ok"])
        r = self.client.post(reverse("operating:order_pack_reserve_add", args=[self.laleli_half.pk]),
                             {"barcode": "ERG1", "pack_id": sack.pk, "place_only": "1"})
        self.assertTrue(r.json()["ok"], r.json())
        self.assertEqual(set(sack.stock_reservations.values_list("pk", flat=True)),
                         {self.laleli_roll.pk, self.ergene_roll.pk})
        # The scan placed the roll; it did not take it onto the Laleli half.
        self.ergene_roll.refresh_from_db()
        self.assertEqual(self.ergene_roll.order_id, self.ergene_half.pk)

        page = self.client.get(reverse("operating:order_packing_list", args=[self.ergene_half.pk]))
        self.assertEqual(page.context["number"], "ORD-2026-000023 & ORD-2026-000024")
        self.assertEqual([[row["barcode"] for row in p.rows] for p in page.context["packs"]],
                         [["LAL1", "ERG1"]])
        self.assertEqual(page.context["unassigned_count"], 0)

    def test_the_first_roll_in_a_sack_starts_packing_both_halves(self, _upload):
        sack = Pack.objects.create(order=self.laleli_half, pack_number=1)
        # Taking a roll out of every sack is not packing.
        self.client.post(reverse("operating:order_pack_reserve_assign_pack",
                                 args=[self.ergene_half.pk]),
                         {"reservation_id": self.laleli_roll.pk, "pack_id": ""})
        self.laleli_half.refresh_from_db()
        self.assertEqual(self.laleli_half.order_status, "pending")

        self.client.post(reverse("operating:order_pack_reserve_add", args=[self.ergene_half.pk]),
                         {"barcode": "ERG1", "pack_id": sack.pk, "place_only": "1"})
        for half in (self.laleli_half, self.ergene_half):
            half.refresh_from_db()
            self.assertEqual(half.order_status, "packaging")

        # A second sack with the Laleli roll; emptying the first is not
        # emptying the order.
        Pack.objects.filter(pk=sack.pk).update(pack_number=2)
        other = Pack.objects.create(order=self.laleli_half, pack_number=1)
        self.client.post(reverse("operating:order_pack_reserve_assign_pack",
                                 args=[self.laleli_half.pk]),
                         {"reservation_id": self.laleli_roll.pk, "pack_id": other.pk})
        self.client.post(reverse("operating:order_pack_reserve_assign_pack",
                                 args=[self.laleli_half.pk]),
                         {"reservation_id": self.ergene_roll.pk, "pack_id": ""})
        self.laleli_half.refresh_from_db()
        self.assertEqual(self.laleli_half.order_status, "packaging")

        # Deleting the last sack with goods in it leaves nothing packed.
        self._crud(self.laleli_half, action="delete_pack", pack_id=other.pk)
        for half in (self.laleli_half, self.ergene_half):
            half.refresh_from_db()
            self.assertEqual(half.order_status, "pending")

    def test_unpacking_the_last_roll_opens_the_order_again(self, _upload):
        sack = Pack.objects.create(order=self.laleli_half, pack_number=1)
        url = reverse("operating:order_pack_reserve_assign_pack", args=[self.laleli_half.pk])
        for roll in (self.laleli_roll, self.ergene_roll):
            self.client.post(url, {"reservation_id": roll.pk, "pack_id": sack.pk})
        self.client.post(url, {"reservation_id": self.laleli_roll.pk, "pack_id": ""})
        self.laleli_half.refresh_from_db()
        self.assertEqual(self.laleli_half.order_status, "packaging")
        self.client.post(reverse("operating:order_pack_reserve_remove", args=[self.ergene_half.pk]),
                         {"reservation_id": self.ergene_roll.pk})
        for half in (self.laleli_half, self.ergene_half):
            half.refresh_from_db()
            self.assertEqual(half.order_status, "pending")

    def test_the_customer_hears_of_packing_once(self, _upload):
        Order.objects.filter(pk__in=[self.laleli_half.pk, self.ergene_half.pk]).update(notify_customer=True)
        sack = Pack.objects.create(order=self.laleli_half, pack_number=1)
        url = reverse("operating:order_pack_reserve_assign_pack", args=[self.laleli_half.pk])
        with patch("operating.order_notifications.send_order_event_email") as send:
            for pack_id in (sack.pk, "", sack.pk):
                with self.captureOnCommitCallbacks(execute=True):
                    self.client.post(url, {"reservation_id": self.laleli_roll.pk, "pack_id": pack_id})
        self.laleli_half.refresh_from_db()
        self.assertEqual(self.laleli_half.order_status, "packaging")
        self.assertEqual([c.args[1] for c in send.call_args_list], ["packaging"])

    def test_a_roll_off_another_order_is_still_refused(self, _upload):
        sack = Pack.objects.create(order=self.laleli_half, pack_number=1)
        stranger = Order.objects.create(order_number="ORD-2026-000099")
        OrderStockReservation.objects.filter(pk=self.ergene_roll.pk).update(order=stranger)
        r = self.client.post(reverse("operating:order_pack_reserve_assign_pack",
                                     args=[self.laleli_half.pk]),
                             {"reservation_id": self.ergene_roll.pk, "pack_id": sack.pk})
        self.assertEqual(r.status_code, 404)
        r = self.client.post(reverse("operating:order_pack_reserve_add", args=[self.laleli_half.pk]),
                             {"barcode": "ERG1", "pack_id": sack.pk, "place_only": "1"})
        self.assertFalse(r.json()["ok"])

    def test_packages_are_numbered_as_one_run_across_the_halves(self, _upload):
        # An older split, packed before: a package #1 on each half.
        Pack.objects.create(order=self.laleli_half, pack_number=1)
        theirs = Pack.objects.create(order=self.ergene_half, pack_number=1)
        r = self._crud(self.ergene_half, action="add_pack")
        added = Pack.objects.get(pk=r.json()["pack"]["id"])
        self.assertEqual((added.order_id, added.pack_number), (self.laleli_half.pk, 2))

        # Deleting one closes the gap over both halves; each keeps its own
        # half's number in its code.
        r = self._crud(self.laleli_half, action="delete_pack", pack_id=added.pk)
        self.assertTrue(r.json()["ok"])
        theirs.refresh_from_db()
        self.assertEqual((theirs.pack_number, theirs.code), (2, "PK-2026-000024-2"))

    def test_the_documents_name_both_halves(self, _upload):
        sack = Pack.objects.create(order=self.laleli_half, pack_number=1)
        OrderStockReservation.objects.filter(
            pk__in=[self.laleli_roll.pk, self.ergene_roll.pk]).update(pack=sack)
        with patch("operating.pack_label.labels_pdf_bytes", return_value=b"%PDF") as labels:
            self.client.get(reverse("operating:order_pack_labels", args=[self.ergene_half.pk]))
        self.assertEqual(labels.call_args.args[0], [sack])
        xlsx = self.client.get(reverse("operating:export_packing_list_excel",
                                       args=[self.ergene_half.pk]))
        import io
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
        self.assertEqual(ws["B3"].value, "ORD-2026-000023 & ORD-2026-000024")
        self.assertEqual([ws.cell(row=r, column=7).value for r in (8, 9)], ["LAL1", "ERG1"])
        # And the Ergene half's page offers the list, though the sack is
        # on the Laleli half.
        page = self.client.get(reverse("operating:order_detail", args=[self.ergene_half.pk]))
        self.assertContains(page, reverse("operating:order_packing_list", args=[self.ergene_half.pk]))
