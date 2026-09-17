"""A sales rep reads the warehouse pages without ever reading a cost.

Where staff see a unit cost, a rep sees the sales price derived from it
(erp.roles.sales_rep_price); where staff see a figure BUILT from cost —
net worth, a line total, the stock value — a rep sees nothing at all.
The assertions look for the numbers in the raw response, so a value
left in a title= or data- attribute fails them just as a visible one would.
"""
import re
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from accounting.models import Book, CurrencyCategory
from authentication.models import Permission
from operating.models import Warehouse, WarehouseProduct, WarehouseProductItem

# 40 m at 7.3319: worth 293.276, and priced for a rep at
# 7.3319 × 1.10 = 8.065… → 8.10.
UNIT_COST = Decimal("7.3319")
METRES = Decimal("40")
PURCHASE_PRICE = Decimal("6.1234")


def _num(text):
    """A pattern matching `text` with either decimal separator."""
    return re.compile(re.escape(text).replace(r"\.", "[.,]"))


class WarehouseCostIsHiddenFromSalesRep(TestCase):
    @patch("marketing.utils.bunny_storage.upload_to_bunny")
    def setUp(self, mock_upload):
        mock_upload.return_value = "https://mock-cdn.net/qr.png"
        usd = CurrencyCategory.objects.create(code="USD", name="US Dollar", symbol="$")
        book = Book.objects.create(name="Laleli Fabric", base_currency=usd)
        self.warehouse = Warehouse.objects.create(name="Laleli depo", accounting_book=book)
        self.product = WarehouseProduct.objects.create(
            warehouse=self.warehouse, name="GREK TÜL", sku="K24593.G07",
            quantity=METRES, purchase_price=PURCHASE_PRICE, purchase_currency="TRY")
        WarehouseProductItem.objects.create(
            product=self.product, quantity=METRES, quantity_remaining=METRES,
            barcode="ROLL-0001", status="in_stock", unit_cost_base=UNIT_COST)

        self.user = User.objects.create_user("rep_wh_test", password="pw")
        self.user.member.books.add(book)
        self.user.member.default_book = book
        self.user.member.save()
        self.client.force_login(self.user)

    def _make_rep(self):
        perm, _ = Permission.objects.get_or_create(name="sales_rep")
        self.user.member.permissions.add(perm)

    def _get(self, url):
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200, url)
        return resp.content.decode()

    def assertHas(self, html, text):
        self.assertIsNotNone(_num(text).search(html), f"{text!r} not in the page")

    def assertLacks(self, html, text):
        # Reports where the figure turned up, not the whole page.
        found = _num(text).search(html)
        if found:
            self.fail(f"{text!r} in the page: …{html[found.start() - 200:found.end() + 80]}…")

    # --- urls ------------------------------------------------------
    @property
    def list_url(self):
        return reverse("operating:warehouse_list")

    @property
    def detail_url(self):
        return reverse("operating:warehouse_detail", args=[self.warehouse.pk])

    @property
    def product_url(self):
        return reverse("operating:warehouse_product_detail",
                       args=[self.warehouse.pk, self.product.pk])

    @property
    def rolls_url(self):
        return reverse("operating:warehouse_product_rolls",
                       args=[self.warehouse.pk, self.product.pk])

    @property
    def group_url(self):
        return reverse("operating:warehouse_group_variants", args=[self.warehouse.pk])

    @property
    def excel_url(self):
        return reverse("operating:warehouse_excel", args=[self.warehouse.pk])

    # --- staff -----------------------------------------------------
    def test_staff_still_see_cost_and_totals(self):
        html = self._get(self.list_url)
        self.assertIn('<span class="num">Net worth</span>', html)
        self.assertIn("$293", html)

        html = self._get(self.detail_url + "?view=variants")
        self.assertIn("Unit cost", html)
        self.assertIn("Net worth (USD)", html)
        self.assertHas(html, "7.3319")
        self.assertHas(html, "293.28")

        html = self._get(self.product_url)
        self.assertIn("Purchase price", html)
        self.assertHas(html, "6.12")
        self.assertHas(html, "7.3319")
        self.assertHas(html, "293.28")
        self.assertIn('name="purchase_price"', html)

        self.assertHas(self._get(self.rolls_url), "7.33")

    def test_staff_excel_keeps_cost_and_total(self):
        values = self._excel_values()
        self.assertIn("Br. Maliyet", values)
        self.assertIn("Toplam (USD)", values)
        self.assertTrue(any("7.3319" in v for v in values), values)

    # --- sales rep -------------------------------------------------
    def test_warehouse_list_shows_a_rep_no_net_worth(self):
        self._make_rep()
        html = self._get(self.list_url)
        self.assertNotIn('<span class="num">Net worth</span>', html)
        self.assertLacks(html, "$293")

    def test_warehouse_detail_shows_a_rep_price_and_no_totals(self):
        self._make_rep()
        for view in ("variants", "groups"):
            with self.subTest(view=view):
                html = self._get(self.detail_url + "?view=" + view)
                self.assertNotIn("Unit cost", html)
                self.assertNotIn("Net worth (USD)", html)
                self.assertLacks(html, "7.33")
                self.assertLacks(html, "293.2")
                self.assertHas(html, "8.10")

    def test_htmx_rows_match_the_rep_header(self):
        self._make_rep()
        resp = self.client.get(self.detail_url + "?view=variants",
                               HTTP_HX_REQUEST="true")
        html = resp.content.decode()
        self.assertEqual(resp.status_code, 200)
        self.assertLacks(html, "7.33")
        self.assertLacks(html, "293.2")
        self.assertHas(html, "8.10")
        # One cell fewer than staff: no Total.
        self.assertNotIn("data-label=\"Total\"", html.split("wh-vrow", 1)[1].split("</tr>")[0])

    def test_group_variant_rows_show_a_rep_price(self):
        self._make_rep()
        # K24593.G07 groups under the part of its SKU before the dot.
        html = self._get(self.group_url + "?base=K24593")
        self.assertLacks(html, "7.33")
        self.assertLacks(html, "293.2")
        self.assertHas(html, "8.10")

    def test_roll_rows_show_a_rep_price(self):
        self._make_rep()
        html = self._get(self.rolls_url)
        self.assertLacks(html, "7.33")
        self.assertHas(html, "8.10")

    def test_product_detail_shows_a_rep_price_only(self):
        self._make_rep()
        html = self._get(self.product_url)
        self.assertNotIn("Purchase price", html)
        self.assertNotIn("Unit cost (USD)", html)
        self.assertNotIn("Total value (USD)", html)
        self.assertNotIn('name="purchase_price"', html)
        self.assertLacks(html, "6.12")
        self.assertLacks(html, "7.33")
        self.assertLacks(html, "293.2")
        self.assertHas(html, "8.10")

    def test_excel_gives_a_rep_price_and_no_total(self):
        self._make_rep()
        values = self._excel_values()
        self.assertNotIn("Br. Maliyet", values)
        self.assertNotIn("Toplam (USD)", values)
        self.assertIn("Br. Fiyat", values)
        self.assertFalse(any("7.33" in v or "293" in v for v in values), values)
        self.assertTrue(any("8.10" in v for v in values), values)

    def _excel_values(self):
        from openpyxl import load_workbook

        resp = self.client.get(self.excel_url)
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(BytesIO(resp.content)).active
        return [str(c.value) for row in ws.iter_rows() for c in row
                if c.value is not None]
