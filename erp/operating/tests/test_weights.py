"""Weights for export papers: a roll's own kilos when its label gave them,
else the fabric's GSM over its width, else nothing; the sack's weight off
the scale; and the packing list's net and gross totals."""
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from accounting.models import Book
from marketing.models import (Product, ProductCategory, ProductVariant,
                              ProductVariantAttribute,
                              ProductVariantAttributeValue)
from operating.models import (Order, OrderItem, OrderStockReservation, Pack,
                              Warehouse, WarehouseProduct, WarehouseProductItem)
from operating.views import _order_weights, _pack_roll_rows, _pack_weights
from operating.views_warehouse import _add_stock_to_variant
from operating.weights import net_kg


def _value(attr_name, value):
    attr, _ = ProductVariantAttribute.objects.get_or_create(name=attr_name)
    row, _ = ProductVariantAttributeValue.objects.get_or_create(
        product_variant_attribute=attr, product_variant_attribute_value=value)
    return row


class NetKg(TestCase):
    def setUp(self):
        self.product = Product.objects.create(title="Velvet", sku="VEL", unit="m",
                                              gsm=Decimal("320"))
        self.variant = ProductVariant.objects.create(product=self.product,
                                                     variant_sku="VEL.RED-280")
        self.variant.product_variant_attribute_values.add(_value("width", "280"))
        wh = Warehouse.objects.create(
            name="WH", accounting_book=Book.objects.get_or_create(name="Laleli")[0])
        self.wp = WarehouseProduct.objects.create(warehouse=wh, name="Velvet", sku="VEL.RED-280")

    def roll(self, weight=None):
        return WarehouseProductItem.objects.create(
            product=self.wp, quantity=Decimal("50"), quantity_remaining=Decimal("50"),
            weight_kg=weight)

    def test_a_weighed_roll_shares_its_weight_by_length(self):
        self.assertEqual(net_kg(Decimal("30"), self.roll(Decimal("45")), self.product, self.variant),
                         (Decimal("27.00"), False))

    def test_an_unweighed_roll_is_worked_out_from_gsm_and_width(self):
        # 0.32 kg/m² × 2.8 m × 30 m
        self.assertEqual(net_kg(Decimal("30"), self.roll(), self.product, self.variant),
                         (Decimal("26.88"), True))

    def test_nothing_known_is_none_not_zero(self):
        self.product.gsm = None
        self.assertEqual(net_kg(Decimal("30"), self.roll(), self.product, self.variant),
                         (None, False))

    def test_no_width_no_estimate(self):
        self.variant.product_variant_attribute_values.clear()
        self.assertEqual(net_kg(Decimal("30"), self.roll(), self.product, self.variant),
                         (None, False))

    def test_goods_counted_in_pieces_are_not_estimated(self):
        self.product.unit = "piece"
        self.assertEqual(net_kg(Decimal("3"), self.roll(), self.product, self.variant),
                         (None, False))


class PackingListWeights(TestCase):
    @patch("marketing.utils.bunny_storage.upload_to_bunny")
    def setUp(self, mock_upload):
        mock_upload.return_value = "https://mock-cdn.net/qr.png"
        self.client.force_login(User.objects.create_superuser("w", "w@t.com", "pw"))
        fabric = ProductCategory.objects.create(name="fabric")
        product = Product.objects.create(title="Velvet", sku="VEL", unit="m",
                                         gsm=Decimal("320"), category=fabric)
        variant = ProductVariant.objects.create(product=product, variant_sku="VEL.RED-280")
        variant.product_variant_attribute_values.add(_value("width", "280"))
        wh = Warehouse.objects.create(
            name="WH", accounting_book=Book.objects.get_or_create(name="Laleli")[0])
        wp = WarehouseProduct.objects.create(warehouse=wh, name="Velvet", sku="VEL.RED-280")
        self.order = Order.objects.create(order_number="DK0000900")
        item = OrderItem.objects.create(order=self.order, product=product,
                                        product_variant=variant, quantity=60, price=10)
        self.pack = Pack.objects.create(order=self.order, pack_number=1)
        for code, weight in (("R1", Decimal("45")), ("R2", None)):
            roll = WarehouseProductItem.objects.create(
                product=wp, quantity=Decimal("50"), quantity_remaining=Decimal("50"),
                barcode=code, weight_kg=weight)
            OrderStockReservation.objects.create(
                order=self.order, order_item=item, stock_item=roll,
                warehouse_product=wp, quantity=Decimal("30"), pack=self.pack)

    def test_rows_and_package_add_up(self):
        rows = _pack_roll_rows(self.pack)
        self.assertEqual([(r["net_kg"], r["net_kg_estimated"]) for r in rows],
                         [(Decimal("27.00"), False), (Decimal("26.88"), True)])
        w = _pack_weights(self.pack, rows)
        self.assertEqual((w["net_kg"], w["net_kg_estimated"], w["gross_kg"]),
                         (Decimal("53.88"), True, None))

    def test_an_order_total_is_none_while_a_sack_is_unweighed(self):
        w = _pack_weights(self.pack, _pack_roll_rows(self.pack))
        self.assertEqual(_order_weights([w]), (Decimal("53.88"), None))
        self.pack.gross_kg = Decimal("55.5")
        w = _pack_weights(self.pack, _pack_roll_rows(self.pack))
        self.assertEqual(_order_weights([w]), (Decimal("53.88"), Decimal("55.5")))

    def test_an_unweighed_empty_sack_is_left_out(self):
        empty = Pack.objects.create(order=self.order, pack_number=2)
        self.pack.gross_kg = Decimal("55.5")
        weights = [_pack_weights(p, _pack_roll_rows(p)) for p in (self.pack, empty)]
        self.assertEqual(_order_weights(weights), (Decimal("53.88"), Decimal("55.5")))

    def test_the_packer_types_the_sack_weight(self):
        url = reverse("operating:order_packing_list", kwargs={"pk": self.order.pk})
        r = self.client.post(url, {"action": "set_gross_kg", "pack_id": self.pack.pk,
                                   "gross_kg": "55,5"})
        self.assertEqual(r.json(), {"ok": True, "gross_kg": "55.50"})
        self.pack.refresh_from_db()
        self.assertEqual(self.pack.gross_kg, Decimal("55.50"))
        r = self.client.post(url, {"action": "set_gross_kg", "pack_id": self.pack.pk,
                                   "gross_kg": ""})
        self.pack.refresh_from_db()
        self.assertIsNone(self.pack.gross_kg)

    def test_nonsense_is_refused(self):
        url = reverse("operating:order_packing_list", kwargs={"pk": self.order.pk})
        for bad in ("heavy", "-3", "0"):
            r = self.client.post(url, {"action": "set_gross_kg", "pack_id": self.pack.pk,
                                       "gross_kg": bad})
            self.assertEqual(r.status_code, 400, bad)

    def test_the_excel_carries_net_and_gross(self):
        Pack.objects.filter(pk=self.pack.pk).update(gross_kg=Decimal("55.5"))
        r = self.client.get(reverse("operating:export_packing_list_excel",
                                    kwargs={"pk": self.order.pk}))
        ws = openpyxl.load_workbook(BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[7]][8:10], ["Net kg", "Gross kg"])
        self.assertEqual([ws.cell(row=8, column=9).value, ws.cell(row=9, column=9).value],
                         [27.0, 26.88])
        self.assertEqual(ws.cell(row=8, column=10).value, 55.5)
        totals = {ws.cell(row=r, column=9).value: ws.cell(row=r, column=10).value
                  for r in range(10, ws.max_row + 1)}
        self.assertEqual((totals["Total Net kg"], totals["Total Gross kg"]), (53.88, 55.5))

    def test_the_page_and_pdf_render(self):
        page = self.client.get(reverse("operating:order_packing_list",
                                       kwargs={"pk": self.order.pk}))
        self.assertContains(page, "~26.88 kg")
        pdf = self.client.get(reverse("operating:order_packing_list_pdf",
                                      kwargs={"pk": self.order.pk}))
        self.assertTrue(pdf.content.startswith(b"%PDF"))


class ReceiptRollWeight(TestCase):
    def test_a_roll_keeps_the_weight_its_label_gave(self):
        wh = Warehouse.objects.create(
            name="WH", accounting_book=Book.objects.get_or_create(name="Laleli")[0])
        wp = WarehouseProduct.objects.create(warehouse=wh, name="Velvet", sku="VEL.RED-280")
        codes = iter(["A1", "A2"])
        _added, ids = _add_stock_to_variant(
            wp, [{"qty": 50, "weight_kg": "45,2"}, {"qty": 40, "weight_kg": None}],
            lambda: next(codes), None)
        self.assertEqual(
            list(WarehouseProductItem.objects.filter(pk__in=ids).order_by("pk")
                 .values_list("weight_kg", flat=True)),
            [Decimal("45.20"), None])
